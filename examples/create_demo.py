"""Create small, synthetic workbooks for exercising the UI."""
import sys
from pathlib import Path

import openpyxl

# Allow direct execution from any working directory.
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from gmcalc.core import BILL_PREFIX, FIELDS, revenue_headers


def create(destination):
    destination.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Q3 Baseline Data'
    for c, label in {2: 'Segment', 3: 'LoB', 4: 'Service Line',
                     **dict(zip([7, 10, 18, 21], revenue_headers('Q3')))}.items():
        ws.cell(1, c, label)
    for r, values in enumerate([
        ['Banking', 'Account A', 'Consulting', 1, 2, 4, 5],
        ['Banking', 'Account A', 'Technology', 2, -1, 6, 5],
        ['Insurance', 'MetLife US', 'Technology', 1, 2, 3, 4],
    ], 2):
        for c, value in zip([2, 3, 4, 7, 10, 18, 21], values):
            ws.cell(r, c, value)
    wb.save(destination / 'demo_baseline.xlsx')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws.append([f + ' 2026-10-21' if f == BILL_PREFIX else f for f in FIELDS])
    for emp, segment, customer, location, cost in [
        ('0001', 'Banking', 'Account A', 'ONSITE', 10000000),
        ('0002', 'Banking', 'Account A', 'OFFSITE', 10000000),
        ('0003', 'Insurance', 'MetLife US', 'OFFSITE', 0),
    ]:
        rec = dict.fromkeys(FIELDS, 'Example')
        rec.update({'Employee Number': emp, 'Segment Wise': segment,
                    'Customer': customer, 'Project Location wrt India': location, 'Cost': cost})
        ws.append([rec[f] for f in FIELDS])
    wb.save(destination / 'demo_headcount.xlsx')


if __name__ == '__main__':
    create(Path(sys.argv[1] if len(sys.argv) > 1 else '/tmp/gm-demo'))
