"""Optional integration check: run with a local LibreOffice/soffice executable."""
import io
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from examples.create_demo import create
from gmcalc.core import headcount


def main():
    engine = shutil.which('soffice') or shutil.which('libreoffice')
    if not engine:
        raise SystemExit('LibreOffice is required for this optional integration check')
    with tempfile.TemporaryDirectory(prefix='gm-excel-') as tmp:
        root = Path(tmp)
        create(root)
        result = headcount((root / 'demo_headcount.xlsx').read_bytes())
        wb = openpyxl.load_workbook(io.BytesIO(result.workbook))
        wb.save(root / 'unfiltered.xlsx')
        ws = wb['Costs By Segment']
        ws.auto_filter.add_filter_column(0, ['Insurance', 'GRAND TOTAL'])
        ws.row_dimensions[2].hidden = True
        ws.row_dimensions[3].hidden = True
        wb.save(root / 'filtered.xlsx')
        output = root / 'recalculated'
        output.mkdir()
        subprocess.run([engine, '-env:UserInstallation=' + (root / 'profile').as_uri(),
                        '--headless', '--convert-to', 'xlsx', '--outdir', str(output),
                        str(root / 'unfiltered.xlsx'), str(root / 'filtered.xlsx')],
                       check=True, timeout=60)
        for name, expected in [('unfiltered', 20000000), ('filtered', 0)]:
            ws = openpyxl.load_workbook(output / (name + '.xlsx'), data_only=True)['Costs By Segment']
            assert ws['E3'].value == expected, (name, 'segment subtotal', ws['E3'].value)
            assert ws['E6'].value == expected, (name, 'grand total', ws['E6'].value)
    print('PASS: LibreOffice nested SUBTOTAL recalculation, unfiltered and filtered/hidden rows')


if __name__ == '__main__':
    main()
