from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from numbers import Real

import openpyxl
import xlsxwriter
from openpyxl.styles import Font
from openpyxl.packaging.custom import StringProperty

TOLERANCE = Decimal('0.000001')
BILL_PREFIX = 'Billable / Non-billable - As on'
FIELDS = ['Employee Number', 'Employee Grade', 'BPS/Non BPS', 'WON/SWON',
          'Project Number', 'Cost', 'SP Type', 'Project Sub-SP', 'Group Customer',
          'Customer', 'BFSI Account Name', 'Project Location wrt India',
          'Project Type', 'LOB', 'Segment Wise', BILL_PREFIX, 'New Mapping',
          'New Mapping Segment', 'NCI Ageing Days', 'Domain', 'Segment']


class ValidationError(ValueError):
    pass


def norm(value):
    return '' if value is None else re.sub(r'\s+', ' ', str(value)).strip()


def key(value):
    # Business keys only trim: internal spaces, spelling and case are significant.
    return '' if value is None else str(value).strip()


def number(value):
    if isinstance(value, bool) or not isinstance(value, (Real, Decimal)):
        raise ValidationError(f'Expected a finite number, received {value!r}')
    result = Decimal(str(value))
    if not result.is_finite():
        raise ValidationError('Non-finite numeric value')
    return result


@dataclass
class Result:
    workbook: bytes
    tables: dict
    issues: list = field(default_factory=list)
    checks: list = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    detail: dict = field(default_factory=dict)

    @property
    def status(self):
        return 'FAIL' if self.issues or any(not c['passed'] for c in self.checks) else 'PASS'

    def report(self):
        return dict(status=self.status, tolerance_inr_crores=str(TOLERANCE),
                    metadata=self.metadata, checks=self.checks, issues=self.issues)


def check(checks, name, actual, expected, tolerance=TOLERANCE):
    delta = number(actual) - number(expected)
    checks.append(dict(name=name, actual=str(actual), expected=str(expected),
                       difference=str(delta), passed=abs(delta) <= tolerance))


def load(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = z.namelist()
            # openpyxl cannot safely round-trip all Excel package extensions.
            unsupported = ('vbaProject', 'externalLinks/', 'activeX/', 'embeddings/',
                           'slicer', 'pivot', 'drawings/', 'vmlDrawing', 'customXml/',
                           'connections.xml', 'queryTables/', 'threadedComments/')
            if any(any(p in n for p in unsupported) for n in names):
                raise ValidationError('Unsupported workbook features (macros, links, drawings, pivots or extensions). Provide a plain .xlsx copy; original is untouched.')
            if any(b'<extLst' in z.read(n) or b':extLst' in z.read(n)
                   for n in names if n.endswith('.xml')):
                raise ValidationError('Workbook contains unsupported Excel extensions. Provide a plain .xlsx copy.')
        return (openpyxl.load_workbook(io.BytesIO(data)),
                openpyxl.load_workbook(io.BytesIO(data), data_only=True))
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValidationError('Invalid .xlsx workbook') from exc


def cell_value(ws, cached, row, col, formula_caches_verified):
    cell = ws.cell(row, col)
    if cell.data_type == 'e':
        raise ValidationError(f'{ws.title}!{cell.coordinate}: Excel error {cell.value}')
    if cell.data_type == 'f':
        if not formula_caches_verified:
            raise ValidationError('Source formulas require recalculation in Excel/LibreOffice and explicit cache-freshness confirmation.')
        value = cached.cell(row, col).value
        if value is None or cached.cell(row, col).data_type == 'e':
            raise ValidationError(f'{ws.title}!{cell.coordinate}: missing/invalid formula cache')
        return value
    return cell.value


def sheet_digest(ws):
    def canonical(value):
        if isinstance(value, (Real, Decimal)) and not isinstance(value, bool):
            return format(number(value).normalize(), 'f')
        return str(value)
    cells = [(c.coordinate, c.data_type, canonical(c.value)) for row in ws for c in row if c.value not in (None, '')]
    return hashlib.sha256(json.dumps(cells, ensure_ascii=False).encode()).hexdigest()


def set_property(wb, name, value):
    if name in wb.custom_doc_props.names:
        del wb.custom_doc_props[name]
    wb.custom_doc_props.append(StringProperty(name=name, value=value))


def generated(wb, name, headers):
    if name in wb.sheetnames:
        if wb[name].sheet_properties.codeName != 'GMcalculationGeneratedV1':
            raise ValidationError(f'{name!r} contains unrecognized content; refusing replacement. Use the original input or remove/rename the sheet in a separate copy.')
        prop = 'GMcalc.digest.' + name
        if prop not in wb.custom_doc_props.names or wb.custom_doc_props[prop].value != sheet_digest(wb[name]):
            raise ValidationError(f'{name!r} was modified after generation; refusing replacement. Preserve your changes in another copy.')
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(headers)
    for cell in ws[1]:
        if isinstance(cell.value, str):
            cell.data_type = 's'
    ws.sheet_properties.codeName = 'GMcalculationGeneratedV1'
    return ws


def style(ws, numeric_start=1):
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(48, max(18, len(str(col[0].value or '')) + 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row[numeric_start:]:
            if isinstance(cell.value, (Real, Decimal)) or cell.data_type == 'f':
                cell.number_format = '#,##0.00'


def save(wb):
    for ws in wb:
        if ws.sheet_properties.codeName == 'GMcalculationGeneratedV1':
            set_property(wb, 'GMcalc.digest.' + ws.title, sheet_digest(ws))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def revenue_headers(quarter):
    if quarter not in ('Q1', 'Q2', 'Q3', 'Q4'):
        raise ValidationError('Quarter must be Q1, Q2, Q3 or Q4')
    return ['Q1 Rev Actual', 'Q2 Rev Actual', f'{quarter} Target(Per AOP) Catch up',
            f'{quarter} Current Rev (MIH) - Total']


def baseline(data, quarter, header_row=1, excluded_rows=(), blanks_are_zero=False,
             formula_caches_verified=False, filename='baseline.xlsx'):
    wb, cached = load(data)
    sheet = f'{quarter} Baseline Data'
    if sheet not in wb.sheetnames:
        raise ValidationError(f'Missing sheet {sheet}')
    ws, cv = wb[sheet], cached[sheet]
    measures = revenue_headers(quarter)
    mapping = {2: 'Segment', 3: 'LoB', 4: 'Service Line', **dict(zip((7, 10, 18, 21), measures))}
    if not 1 <= header_row <= ws.max_row:
        raise ValidationError('Invalid baseline header row')
    for col, label in mapping.items():
        # Merged/grouped headings: verify all labels in the user-selected header band.
        candidates = [norm(ws.cell(r, col).value) for r in range(1, header_row + 1)]
        if norm(label) not in candidates:
            raise ValidationError(f'Column {openpyxl.utils.get_column_letter(col)} must identify {label!r} in the header band ending at row {header_row}')
    exclusions = set(excluded_rows)
    if any(r <= header_row or r > ws.max_row for r in exclusions):
        raise ValidationError('Excluded row outside data range')
    records, issues, checks = [], [], []
    blank_rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        if r in exclusions:
            continue
        if all(ws.cell(r, c).value is None for c in range(1, ws.max_column + 1)):
            blank_rows.append(r)
            continue
        values = {c: cell_value(ws, cv, r, c, formula_caches_verified) for c in mapping}
        dims = []
        for c, label in ((2, 'Segment'), (3, 'LoB'), (4, 'Service Line')):
            value = key(values[c])
            if not value:
                issues.append(f'Row {r}: missing {label}; revenue retained in exception bucket')
                value = f'(Missing {label})'
            if re.search(r'\b(total|subtotal)\b', value, re.I):
                raise ValidationError(f'Row {r}: possible subtotal/total; verify and list exclusion explicitly (or rename a genuine detail label in a separate input copy).')
            dims.append(value)
        amounts = []
        for c in (7, 10, 18, 21):
            if values[c] is None and blanks_are_zero:
                amounts.append(Decimal(0))
            else:
                try:
                    amounts.append(number(values[c]))
                except ValidationError as exc:
                    raise ValidationError(f'{sheet}!{openpyxl.utils.get_column_letter(c)}{r}: {exc}; confirm genuine blanks as zero if applicable') from exc
        records.append((dims, amounts))
    if not records:
        raise ValidationError('No baseline detail records')
    for dim in range(3):
        variants = defaultdict(set)
        for dims, _ in records:
            variants[dims[dim].casefold()].add(dims[dim])
        for names in variants.values():
            if len(names) > 1:
                issues.append(f'Suspected duplicate names retained separately: {sorted(names)}')
    tables, detail = {}, {}
    source = [sum((a[i] for _, a in records), Decimal(0)) for i in range(4)]
    for name, indices, labels in [('Segment View', (0,), ['Segment']),
                                  ('Account View', (0, 1), ['Segment', 'AccountOrLoB']),
                                  ('SP View', (2,), ['Service Line'])]:
        groups = defaultdict(lambda: [Decimal(0)] * 4)
        for dims, amounts in records:
            k = tuple(dims[i] for i in indices)
            groups[k] = [a + b for a, b in zip(groups[k], amounts)]
        rows = [list(k) + groups[k] for k in sorted(groups)]
        out = generated(wb, name, labels + measures)
        for row in rows:
            out.append([float(v) if isinstance(v, Decimal) else v for v in row])
            for cell in out[out.max_row][:len(labels)]:
                cell.data_type = 's'
        out.sheet_properties.pageSetUpPr.fitToPage = True
        out.oddHeader.center.text = 'Revenue: INR Crores'
        style(out, len(labels))
        for i, measure in enumerate(measures):
            check(checks, f'{name}: {measure}', sum((v[i] for v in groups.values()), Decimal(0)), source[i])
        tables[name] = dict(headers=labels + measures, rows=rows)
        detail[name] = groups
    set_property(wb, 'GMcalc.revenue_issues', json.dumps(issues))
    return Result(save(wb), tables, issues, checks,
                  dict(filename=filename, quarter=quarter, processed_rows=len(records),
                       excluded_rows=sorted(exclusions), blank_rows=blank_rows,
                       blanks_confirmed_zero=blanks_are_zero), detail)


def source_fields(ws, header_row):
    if not 1 <= header_row <= ws.max_row:
        raise ValidationError('Invalid headcount header row')
    headers = [norm(c.value) for c in ws[header_row]]
    columns, output_headers = {}, []
    for label in FIELDS:
        matches = [i + 1 for i, h in enumerate(headers)
                   if (h.casefold().startswith(label.casefold()) if label == BILL_PREFIX else h.casefold() == label.casefold())]
        if len(matches) != 1:
            raise ValidationError(f'{label!r}: expected one header, found {len(matches)}')
        columns[label] = matches[0]
        output_headers.append(ws.cell(header_row, matches[0]).value)
    return columns, output_headers


def headcount(data, header_row=1, formula_caches_verified=False, filename='headcount.xlsx'):
    wb, cached = load(data)
    if 'Data' not in wb.sheetnames:
        raise ValidationError('Missing Data sheet')
    ws, cv = wb['Data'], cached['Data']
    cols, headers = source_fields(ws, header_row)
    records, issues, checks = [], [], []
    ids = Counter()
    for r in range(header_row + 1, ws.max_row + 1):
        if all(ws.cell(r, c).value is None for c in range(1, ws.max_column + 1)):
            continue
        rec = {f: cell_value(ws, cv, r, c, formula_caches_verified) for f, c in cols.items()}
        rec['_row'] = r
        rec['_location'] = key(rec['Project Location wrt India']).upper()
        emp = key(rec['Employee Number'])
        if not emp:
            issues.append(f'Row {r}: blank Employee Number')
        else:
            ids[emp] += 1
        if rec['_location'] not in ('ONSITE', 'OFFSITE'):
            issues.append(f'Row {r}: invalid location {rec["Project Location wrt India"]!r}; excluded from HC and cost buckets')
        try:
            rec['_cost'] = number(rec['Cost'])
        except ValidationError:
            rec['_cost'] = None
            issues.append(f'Row {r}: invalid/blank Cost; excluded from cost sums pending review')
        for f in ('Segment Wise', 'Customer', 'LOB', BILL_PREFIX, 'New Mapping', 'WON/SWON', 'SP Type'):
            if not key(rec[f]):
                issues.append(f'Row {r}: blank grouping field {f}; retained as (Blank)')
        records.append(rec)
    if not records:
        raise ValidationError('No headcount records')
    for emp, count in ids.items():
        if count > 1:
            issues.append(f'Duplicate Employee Number {emp!r}: rows {[r["_row"] for r in records if key(r["Employee Number"]) == emp]}')
    kd = generated(wb, 'Key Data', headers)
    for rec in records:
        kd.append([rec[f] for f in FIELDS])
        # Preserve source number formats (notably zero-padded employee identifiers).
        for j, f in enumerate(FIELDS, 1):
            if isinstance(kd.cell(kd.max_row, j).value, str):
                kd.cell(kd.max_row, j).data_type = 's'
            kd.cell(kd.max_row, j).number_format = ws.cell(rec['_row'], cols[f]).number_format
    style(kd, 21)
    check(checks, 'Key Data record count', len(records), len(records), Decimal(0))
    tables = {'Key Data': dict(headers=headers, rows=[[r[f] for f in FIELDS] for r in records])}
    dash = generated(wb, 'Headcount Dashboard', ['Headcount Dashboard'])
    valid_hc = sum(bool(key(r['Employee Number'])) and r['_location'] in ('ONSITE', 'OFFSITE') for r in records)
    for name, fields in [('Segment Summary', ['Segment Wise']),
                         ('Billability Summary', ['Segment Wise', BILL_PREFIX]),
                         ('New Mapping Summary', ['Segment Wise', 'New Mapping']),
                         ('WON/SWON Summary', ['Segment Wise', 'WON/SWON']),
                         ('SP Type Summary', ['Segment Wise', 'SP Type'])]:
        groups = defaultdict(lambda: [0, 0])
        for rec in records:
            k = tuple(key(rec[f]) or '(Blank)' for f in fields)
            counts = groups[k]
            if key(rec['Employee Number']) and rec['_location'] in ('ONSITE', 'OFFSITE'):
                counts[0 if rec['_location'] == 'ONSITE' else 1] += 1
        rows = [list(k) + v + [sum(v)] for k, v in sorted(groups.items())]
        totals = [sum(v[i] for v in groups.values()) for i in range(2)]
        full_headers = [headers[FIELDS.index(f)] for f in fields] + ['Onsite HC', 'Offsite HC', 'Total HC']
        dash.append([name])
        dash.append(full_headers)
        for row in rows:
            dash.append(row)
        dash.append(['GRAND TOTAL'] + [''] * (len(fields) - 1) + totals + [sum(totals)])
        dash.append([])
        tables[name] = dict(headers=full_headers, rows=rows)
        check(checks, name, sum(totals), valid_hc, Decimal(0))
    dash.append(['Invalid location records', sum(r['_location'] not in ('ONSITE', 'OFFSITE') for r in records)])
    for row in dash:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    style(dash)
    dash.auto_filter.ref = None
    detail = {}
    for name, secondary, metlife in [('Costs By Segment', 'Customer', False), ('Costs for MetLife', 'LOB', True)]:
        groups = defaultdict(lambda: [Decimal(0), Decimal(0)])
        eligible = [r for r in records if not metlife or 'metlife' in key(r['Customer']).casefold()]
        for rec in eligible:
            if rec['_location'] not in ('ONSITE', 'OFFSITE') or rec['_cost'] is None:
                continue
            k = (key(rec['Segment Wise']) or '(Blank)', key(rec[secondary]) or '(Blank)')
            groups[k][0 if rec['_location'] == 'ONSITE' else 1] += rec['_cost']
        out = generated(wb, name, ['Segment Wise', secondary, 'ONSITE COST', 'OFFSITE COST', 'TOTAL COST'])
        rows = []
        for segment in sorted({k[0] for k in groups}):
            start = out.max_row + 1
            for k in sorted(k for k in groups if k[0] == segment):
                v = groups[k]
                row = list(k) + v + [sum(v)]
                rows.append(row)
                out.append([float(x) if isinstance(x, Decimal) else x for x in row])
                out.cell(out.max_row, 1).data_type = 's'
                out.cell(out.max_row, 2).data_type = 's'
            end = out.max_row
            out.append([segment, 'SEGMENT TOTAL'] + [f'=SUBTOTAL(109,{c}{start}:{c}{end})' for c in 'CDE'])
            out.cell(out.max_row, 1).data_type = 's'
        end = out.max_row
        out.append(['GRAND TOTAL', ''] + [f'=SUBTOTAL(109,{c}2:{c}{end})' if end >= 2 else '=0' for c in 'CDE'])
        style(out, 2)
        tables[name] = dict(headers=['Segment Wise', secondary, 'ONSITE COST', 'OFFSITE COST', 'TOTAL COST'], rows=rows)
        detail[name] = {k: sum(v) for k, v in groups.items()}
        # Exclude incomplete groups from margin matching even when some costs are valid.
        invalid_keys = {(key(r['Segment Wise']) or '(Blank)', key(r[secondary]) or '(Blank)') for r in eligible
                        if r['_cost'] is None or r['_location'] not in ('ONSITE', 'OFFSITE')
                        or not key(r['Segment Wise']) or not key(r[secondary])}
        if not metlife:
            detail['invalid_cost_keys'] = invalid_keys
        for i, loc in enumerate(('ONSITE', 'OFFSITE')):
            expected = sum((r['_cost'] for r in eligible if r['_location'] == loc and r['_cost'] is not None), Decimal(0))
            check(checks, f'{name} {loc}', sum((v[i] for v in groups.values()), Decimal(0)), expected)
        out.oddHeader.center.text = 'Costs: monthly INR'
        issues_for_subtotal = 'Excel SUBTOTAL formulas supplied; interactive filter behavior requires Excel/LibreOffice validation (not executed by openpyxl).'
    label = headers[FIELDS.index(BILL_PREFIX)]
    date_text = norm(label)[len(BILL_PREFIX):].strip(' :-')
    # Only record dates that parse unambiguously; retain original label regardless.
    snapshot = None
    for fmt in ('%Y-%m-%d', '%d-%b-%Y', '%d %b %Y'):
        try:
            snapshot = datetime.strptime(date_text, fmt).date().isoformat()
            break
        except ValueError:
            pass
    set_property(wb, 'GMcalc.invalid_cost_keys', json.dumps(sorted(detail['invalid_cost_keys'])))
    set_property(wb, 'GMcalc.upstream_issues', json.dumps(issues))
    return Result(save(wb), tables, issues, checks,
                  dict(filename=filename, processed_rows=len(records), snapshot_date=snapshot,
                       billability_header=label, formula_validation_limitation=issues_for_subtotal,
                       invalid_location_count=sum(r['_location'] not in ('ONSITE', 'OFFSITE') for r in records),
                       duplicate_employee_ids=[k for k, v in ids.items() if v > 1]), detail)


def margin(revenue: Result, costs: Result):
    quarter = revenue.metadata['quarter']
    accounts, segments = revenue.detail['Account View'], revenue.detail['Segment View']
    cost_map = costs.detail['Costs By Segment']
    invalid = costs.detail.get('invalid_cost_keys', set())
    out = io.BytesIO()
    wb = xlsxwriter.Workbook(out, {'in_memory': True, 'strings_to_formulas': False, 'strings_to_urls': False})
    num = wb.add_format({'num_format': '#,##0.00'})
    pct = wb.add_format({'num_format': '0.00%'})
    bold = wb.add_format({'bold': True})
    headers = ['Account', 'Q1 Actual', 'Q2 Actual', *revenue_headers(quarter)[2:],
               'Monthly manpower Cost (MP)', f'{quarter} MP', f'{quarter} NMP', 'Total Cost',
               f'GM Projection versus {quarter} Target', 'GM Projection versus MIH', 'ISU Target']
    mapping, issues, checks, tables = {}, [], [], {}
    used = set()
    for (segment,) in sorted(segments):
        safe = re.sub(r'[:\\/?*\[\]\x00-\x1f]', '_', 'Margin-' + segment).strip("'")
        candidate = safe[:31]
        if len(safe) > 31 or candidate.casefold() in used:
            digest = hashlib.sha256(segment.encode()).hexdigest()
            for length in range(8, 25):
                candidate = safe[:30 - length] + '-' + digest[:length]
                if candidate.casefold() not in used:
                    break
            else:
                raise ValidationError('Unable to generate unique sheet name')
        used.add(candidate.casefold())
        mapping[segment] = candidate
        ws = wb.add_worksheet(candidate)
        ws.write_row(0, 0, headers, bold)
        ws.write_comment(0, 1, 'B–I: INR Crores; F monthly; G–I quarterly. J–L stored fractions.')
        ws.set_column(0, 0, 32)
        ws.set_column(1, 11, 23)
        ws.freeze_panes(1, 0)
        rows = []
        selected = [(k, v) for k, v in sorted(accounts.items()) if k[0] == segment]
        for index, (k, amounts) in enumerate(selected, 1):
            rowno = index + 1
            f = None
            if k in invalid:
                issues.append(f'{segment} / {k[1]}: incomplete source cost/location records; cost unresolved')
            elif k not in cost_map:
                issues.append(f'{segment} / {k[1]}: no exact cost match; cost unresolved')
            elif segment.startswith('(Missing ') or k[1].startswith('(Missing '):
                issues.append(f'{segment} / {k[1]}: missing revenue key; cost unresolved')
            else:
                f = cost_map[k] / Decimal(10000000)
                check(checks, f'{segment}/{k[1]} monthly cost conversion', f * Decimal(10000000), cost_map[k])
            g = f * 3 if f is not None else None
            j = (amounts[2] - g) / amounts[2] if g is not None and amounts[2] != 0 else None
            kval = (amounts[3] - g) / amounts[3] if g is not None and amounts[3] != 0 else None
            for denom, label in ((amounts[2], 'Target'), (amounts[3], 'MIH')):
                if denom == 0:
                    issues.append(f'{segment}/{k[1]}: zero {label} denominator; GM undefined')
            values = [k[1], *amounts, f, g, Decimal(0), g, j, kval, Decimal(0)]
            ws.write_string(index, 0, k[1])
            for col in range(1, 12):
                v = values[col]
                if v is not None:
                    ws.write_number(index, col, float(v), pct if col >= 9 else num)
            for col, formula, value in [
                (6, f'=IF(ISNUMBER(F{rowno}),F{rowno}*3,"")', g),
                (8, f'=IF(ISNUMBER(G{rowno}),G{rowno}+H{rowno},"")', g),
                (9, f'=IF(AND(ISNUMBER(I{rowno}),ISNUMBER(D{rowno}),D{rowno}<>0),(D{rowno}-I{rowno})/D{rowno},"")', j),
                (10, f'=IF(AND(ISNUMBER(I{rowno}),ISNUMBER(E{rowno}),E{rowno}<>0),(E{rowno}-I{rowno})/E{rowno},"")', kval)]:
                ws.write_formula(index, col, formula, pct if col >= 9 else num, float(value) if value is not None else '')
            rows.append(values)
        ws.autofilter(0, 0, len(rows), 11)
        tables[segment] = dict(headers=headers, rows=rows)
        for i, measure in enumerate(revenue_headers(quarter)):
            actual = sum((v[i] for _, v in selected), Decimal(0))
            check(checks, f'{segment} {measure} vs Segment View', actual, segments[(segment,)][i])
            check(checks, f'{segment} {measure} vs Account View', sum((r[i + 1] for r in rows), Decimal(0)), actual)
    for k in sorted(set(cost_map) - set(accounts)):
        issues.append(f'Cost key without revenue account: {k[0]} / {k[1]}')
    meta = dict(quarter=quarter, revenue_filename=revenue.metadata['filename'],
                headcount_filename=costs.metadata['filename'], snapshot_date=costs.metadata['snapshot_date'],
                segment_sheet_mapping=mapping, matching='Case-sensitive exact values after trimming',
                projection='Monthly INR / 10000000 × 3; NMP and ISU Target default to zero')
    report_ws = wb.add_worksheet('Run Information')
    report_ws.set_column(0, 1, 65)
    info = [('Revenue file', meta['revenue_filename']), ('Headcount file', meta['headcount_filename']),
            ('Quarter', quarter), ('Snapshot date', meta['snapshot_date'] or 'Unverified'),
            ('Units', 'INR Crores; F monthly, G–I quarterly; J–L percentages'),
            ('Projection', meta['projection']), ('Segment', 'Sheet name'), *mapping.items()]
    for i, row in enumerate(info):
        report_ws.write_row(i, 0, row)
    wb.close()
    return Result(out.getvalue(), tables, revenue.issues + costs.issues + issues,
                  revenue.checks + costs.checks + checks, meta)


def margin_from_workbooks(revenue_data, cost_data, quarter, formula_caches_verified=False,
                          revenue_filename='baseline.xlsx', cost_filename='headcount.xlsx'):
    """Run Step 3 independently from refreshed revenue and headcount workbooks."""
    rev_wb, rev_cv = load(revenue_data)
    cost_wb, cost_cv = load(cost_data)
    measures = revenue_headers(quarter)
    rev_detail, checks, issues = {}, [], []
    for sheet, labels in [('Segment View', ['Segment']), ('Account View', ['Segment', 'AccountOrLoB'])]:
        if sheet not in rev_wb.sheetnames:
            raise ValidationError(f'Missing sheet {sheet}')
        ws = rev_wb[sheet]
        expected = labels + measures
        if [norm(ws.cell(1, i).value) for i in range(1, len(expected) + 1)] != expected:
            raise ValidationError(f'{sheet}: headers do not match selected quarter and expected layout')
        groups = {}
        for r in range(2, ws.max_row + 1):
            if all(ws.cell(r, c).value is None for c in range(1, len(expected) + 1)):
                continue
            k = tuple(key(cell_value(ws, rev_cv[sheet], r, c, formula_caches_verified)) for c in range(1, len(labels) + 1))
            if not all(k) or k in groups:
                raise ValidationError(f'{sheet} row {r}: missing or duplicate key {k}')
            groups[k] = [number(cell_value(ws, rev_cv[sheet], r, c, formula_caches_verified))
                         for c in range(len(labels) + 1, len(expected) + 1)]
        if not groups:
            raise ValidationError(f'{sheet}: no detail records')
        rev_detail[sheet] = groups
    for segment, account in rev_detail['Account View']:
        if (segment,) not in rev_detail['Segment View']:
            raise ValidationError(f'Account View: unmatched segment {segment!r}')
    if 'Costs By Segment' not in cost_wb.sheetnames:
        raise ValidationError('Missing Costs By Segment sheet')
    ws = cost_wb['Costs By Segment']
    expected = ['Segment Wise', 'Customer', 'ONSITE COST', 'OFFSITE COST', 'TOTAL COST']
    if [norm(ws.cell(1, c).value) for c in range(1, 6)] != expected:
        raise ValidationError('Costs By Segment: unexpected headers')
    cost_map, invalid = {}, set()
    for r in range(2, ws.max_row + 1):
        k = tuple(key(ws.cell(r, c).value) for c in (1, 2))
        if k[0] == 'GRAND TOTAL' or k[1] == 'SEGMENT TOTAL' or not any(k):
            continue
        if not all(k):
            invalid.add(k)
            issues.append(f'Costs By Segment row {r}: missing key')
        if k in cost_map:
            invalid.add(k)
            issues.append(f'Costs By Segment row {r}: ambiguous duplicate key {k}')
        try:
            values = [number(cell_value(ws, cost_cv[ws.title], r, c, formula_caches_verified)) for c in (3, 4, 5)]
            check(checks, f'Cost row {r} total arithmetic (INR)', values[0] + values[1], values[2])
            if abs(values[0] + values[1] - values[2]) > TOLERANCE:
                invalid.add(k)
            cost_map[k] = values[2]
        except ValidationError as exc:
            issues.append(f'Costs By Segment row {r}: {exc}')
            invalid.add(k)
    if 'GMcalc.invalid_cost_keys' in cost_wb.custom_doc_props.names:
        invalid.update(tuple(k) for k in json.loads(cost_wb.custom_doc_props['GMcalc.invalid_cost_keys'].value))
    if 'GMcalc.upstream_issues' in cost_wb.custom_doc_props.names:
        issues.extend(json.loads(cost_wb.custom_doc_props['GMcalc.upstream_issues'].value))
    snapshot = None
    if 'Data' in cost_wb.sheetnames:
        # Search header metadata only; never infer a date from file timestamps.
        for row in cost_wb['Data'].iter_rows(max_row=min(30, cost_wb['Data'].max_row)):
            for cell in row:
                label = norm(cell.value)
                if label.casefold().startswith(BILL_PREFIX.casefold()):
                    text = label[len(BILL_PREFIX):].strip(' :-')
                    try:
                        snapshot = datetime.strptime(text, '%Y-%m-%d').date().isoformat()
                    except ValueError:
                        pass
    rev_issues = json.loads(rev_wb.custom_doc_props['GMcalc.revenue_issues'].value) if 'GMcalc.revenue_issues' in rev_wb.custom_doc_props.names else []
    rev = Result(b'', {}, rev_issues, [], dict(quarter=quarter, filename=revenue_filename), rev_detail)
    costs = Result(b'', {}, issues, checks, dict(filename=cost_filename, snapshot_date=snapshot),
                   {'Costs By Segment': cost_map, 'invalid_cost_keys': invalid})
    return margin(rev, costs)
