import io
from decimal import Decimal

import openpyxl
import pytest
import xlsxwriter

from gmcalc.core import (BILL_PREFIX, FIELDS, ValidationError, baseline,
                         headcount, margin, revenue_headers, save)


def baseline_input(rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Q3 Baseline Data'
    for c, label in {2: 'Segment', 3: 'LoB', 4: 'Service Line',
                     **dict(zip([7, 10, 18, 21], revenue_headers('Q3')))}.items():
        ws.cell(1, c, label)
    for i, (segment, account, service, amounts) in enumerate(rows or [
        ('Banking', 'Account A', 'SP1', [1, 2, 4, 5]),
        ('Banking', 'Account A', 'SP2', [2, -1, 6, 5]),
        ('Banking', 'Missing cost', 'SP1', [1, 1, 0, 0]),
        ('Insurance', 'MetLife US', 'SP2', [1, 2, 3, 4]),
    ], 2):
        for c, value in zip([2, 3, 4, 7, 10, 18, 21], [segment, account, service, *amounts]):
            ws.cell(i, c, value)
    wb.create_sheet('Unrelated')['A1'] = 'keep me'
    return save(wb)


def hc_input(records=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    # Deliberately reverse columns: extraction must be label based.
    fields = list(reversed(FIELDS))
    ws.append([BILL_PREFIX + ' 2026-10-21' if f == BILL_PREFIX else f for f in fields])
    for record in records or [
        {'Employee Number': '0001', 'Segment Wise': 'Banking', 'Customer': 'Account A', 'Cost': 10000000, 'Project Location wrt India': ' onsite '},
        {'Employee Number': '0002', 'Segment Wise': 'Banking', 'Customer': 'Account A', 'Cost': 10000000, 'Project Location wrt India': 'OFFSITE'},
        {'Employee Number': '0003', 'Segment Wise': 'Insurance', 'Customer': 'MetLife US', 'Cost': 0, 'Project Location wrt India': 'OFFSITE'},
    ]:
        values = {f: 'Default' for f in FIELDS}
        values.update(record)
        ws.append([values[f] for f in fields])
    return save(wb)


def test_end_to_end_units_reconciliation_and_formula_caches():
    rev = baseline(baseline_input(), 'Q3')
    hc = headcount(hc_input())
    assert rev.status == hc.status == 'PASS'
    assert hc.metadata['snapshot_date'] == '2026-10-21'
    assert hc.tables['Key Data']['rows'][0][0] == '0001'
    assert hc.detail['Costs By Segment'][('Banking', 'Account A')] == 20000000
    result = margin(rev, hc)
    banking = result.tables['Banking']['rows']
    assert banking[0][5:11] == [Decimal(2), Decimal(6), Decimal(0), Decimal(6), Decimal('.4'), Decimal('.4')]
    assert banking[1][5] is None
    assert banking[1][9] is None
    assert result.tables['Insurance']['rows'][0][5] == 0
    assert all(c['passed'] for c in result.checks)
    assert result.status == 'FAIL'  # unmatched account and zero denominators
    formula_wb = openpyxl.load_workbook(io.BytesIO(result.workbook))
    cache_wb = openpyxl.load_workbook(io.BytesIO(result.workbook), data_only=True)
    assert formula_wb['Margin-Banking']['G2'].value == '=IF(ISNUMBER(F2),F2*3,"")'
    assert cache_wb['Margin-Banking']['G2'].value == 6
    assert cache_wb['Margin-Banking']['J2'].value == .4
    assert cache_wb['Margin-Banking']['F3'].value is None
    assert cache_wb['Margin-Banking']['G3'].value is None
    assert formula_wb['Margin-Banking']['J2'].number_format == '0.00%'
    assert openpyxl.load_workbook(io.BytesIO(rev.workbook))['Unrelated']['A1'].value == 'keep me'


def test_cost_subtotals_are_nested_filter_aware_formulas():
    result = headcount(hc_input())
    ws = openpyxl.load_workbook(io.BytesIO(result.workbook))['Costs By Segment']
    assert ws['C3'].value == '=SUBTOTAL(109,C2:C2)'
    assert ws['C6'].value == '=SUBTOTAL(109,C2:C5)'
    assert ws.freeze_panes == 'A2'
    assert ws.auto_filter.ref == 'A1:E6'
    assert 'not executed' in result.metadata['formula_validation_limitation']


def test_invalid_cost_and_location_do_not_create_false_zero_margin():
    hc = headcount(hc_input([
        {'Employee Number': 'x', 'Segment Wise': 'Banking', 'Customer': 'Account A', 'Cost': None, 'Project Location wrt India': 'ONSITE'},
        {'Employee Number': 'x', 'Segment Wise': 'Banking', 'Customer': 'Account A', 'Cost': 10, 'Project Location wrt India': 'OFFSITE'},
        {'Employee Number': None, 'Segment Wise': 'Insurance', 'Customer': 'MetLife US', 'Cost': 0, 'Project Location wrt India': 'Unknown'},
    ]))
    assert hc.status == 'FAIL'
    assert hc.metadata['invalid_location_count'] == 1
    assert hc.metadata['duplicate_employee_ids'] == ['x']
    result = margin(baseline(baseline_input(), 'Q3'), hc)
    assert result.tables['Banking']['rows'][0][5] is None
    assert result.tables['Insurance']['rows'][0][5] is None


def test_blank_and_text_revenue_require_action():
    source = baseline_input([('A', 'B', 'C', [None, 0, 0, 0])])
    with pytest.raises(ValidationError, match='confirm genuine blanks'):
        baseline(source, 'Q3')
    assert baseline(source, 'Q3', blanks_are_zero=True).status == 'PASS'
    with pytest.raises(ValidationError):
        baseline(baseline_input([('A', 'B', 'C', ['1', 0, 0, 0])]), 'Q3', blanks_are_zero=True)


def test_missing_groups_retained_negative_revenue_and_subtotal_exclusion():
    source = baseline_input([('', 'B', 'C', [-2, 1, 1, 1]), ('Grand Total', '', '', [-2, 1, 1, 1])])
    with pytest.raises(ValidationError, match='possible subtotal'):
        baseline(source, 'Q3')
    result = baseline(source, 'Q3', excluded_rows=[3])
    assert result.detail['Segment View'][('(Missing Segment)',)][0] == -2
    assert result.status == 'FAIL'
    assert all(c['passed'] for c in result.checks)


def test_header_identity_and_manual_output_protection():
    wb = openpyxl.load_workbook(io.BytesIO(baseline_input()))
    wb['Q3 Baseline Data']['U1'] = 'Wrong meaning'
    with pytest.raises(ValidationError, match='Column U'):
        baseline(save(wb), 'Q3')
    wb = openpyxl.load_workbook(io.BytesIO(baseline_input()))
    wb.create_sheet('Segment View')['A1'] = 'Manual content'
    with pytest.raises(ValidationError, match='refusing replacement'):
        baseline(save(wb), 'Q3')
    first = baseline(baseline_input(), 'Q3')
    second = baseline(first.workbook, 'Q3')
    assert first.tables == second.tables


def test_ambiguous_headcount_headers_and_undated_snapshot():
    wb = openpyxl.load_workbook(io.BytesIO(hc_input()))
    ws = wb['Data']
    ws.cell(1, 22, ' Employee Number ')
    with pytest.raises(ValidationError, match='found 2'):
        headcount(save(wb))
    ws.delete_cols(22)
    for cell in ws[1]:
        if str(cell.value).startswith(BILL_PREFIX):
            cell.value = BILL_PREFIX + ' 01/02/2026'
    assert headcount(save(wb)).metadata['snapshot_date'] is None


def test_formula_cache_verification_and_missing_cache():
    wb = openpyxl.load_workbook(io.BytesIO(baseline_input()))
    wb['Q3 Baseline Data']['G2'] = '=1+1'
    with pytest.raises(ValidationError, match='explicit cache'):
        baseline(save(wb), 'Q3')
    with pytest.raises(ValidationError, match='missing/invalid formula cache'):
        baseline(save(wb), 'Q3', formula_caches_verified=True)


def test_segment_sheet_names_and_exact_cross_segment_match():
    names = ['A' * 40, 'A' * 39 + 'B', 'Bad/Name', 'bad/name']
    rev = baseline(baseline_input([(s, 'Account A', 'SP', [1, 1, 1, 1]) for s in names]), 'Q3')
    result = margin(rev, headcount(hc_input()))
    mapping = result.metadata['segment_sheet_mapping']
    assert len({v.casefold() for v in mapping.values()}) == 4
    assert all(len(v) <= 31 and '/' not in v for v in mapping.values())
    assert all(t['rows'][0][5] is None for t in result.tables.values())
    assert mapping == margin(rev, headcount(hc_input())).metadata['segment_sheet_mapping']


def test_metlife_substring_and_blank_categories():
    hc = headcount(hc_input([{'Employee Number': '1', 'Segment Wise': None,
                            'Customer': 'METLIFE US', 'LOB': None,
                            'Cost': -100, 'Project Location wrt India': 'offsite'}]))
    assert hc.detail['Costs for MetLife'][('(Blank)', '(Blank)')] == -100
    assert hc.tables['Segment Summary']['rows'] == [['(Blank)', 0, 1, 1]]
    assert all(c['passed'] for c in hc.checks)


def test_independent_step_three_and_duplicate_cost_matches():
    from gmcalc.core import margin_from_workbooks
    rev = baseline(baseline_input(), 'Q3')
    hc = headcount(hc_input())
    result = margin_from_workbooks(rev.workbook, hc.workbook, 'Q3')
    assert result.tables == margin(rev, hc).tables
    wb = openpyxl.load_workbook(io.BytesIO(hc.workbook))
    wb['Costs By Segment'].append(['Banking', 'Account A', 0, 0, 0])
    result = margin_from_workbooks(rev.workbook, save(wb), 'Q3')
    assert result.tables['Banking']['rows'][0][5] is None
    assert any('duplicate' in i for i in result.issues)


def test_streamlit_ui_loads_and_missing_input_is_reported():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file('app.py').run()
    assert not app.exception
    app.button[0].click().run()
    assert not app.exception
    assert 'Upload a quarterly baseline workbook' in app.error[0].value
    app.radio[0].set_value('Step 2: Headcount').run()
    app.button[0].click().run()
    assert 'Upload a daily headcount workbook' in app.error[0].value


def test_manual_edits_to_generated_output_are_protected():
    result = baseline(baseline_input(), 'Q3')
    wb = openpyxl.load_workbook(io.BytesIO(result.workbook))
    wb['Segment View']['G20'] = 'user note'
    raw = io.BytesIO()
    wb.save(raw)  # External edit: do not refresh the application's digest.
    with pytest.raises(ValidationError, match='modified after generation'):
        baseline(raw.getvalue(), 'Q3')


def test_independent_step_three_preserves_invalid_cost_exceptions():
    from gmcalc.core import margin_from_workbooks
    hc = headcount(hc_input([
        {'Employee Number': '1', 'Segment Wise': 'Banking', 'Customer': 'Account A', 'Cost': 5, 'Project Location wrt India': 'ONSITE'},
        {'Employee Number': '2', 'Segment Wise': 'Banking', 'Customer': 'Account A', 'Cost': None, 'Project Location wrt India': 'ONSITE'},
    ]))
    rev = baseline(baseline_input(), 'Q3')
    result = margin_from_workbooks(rev.workbook, hc.workbook, 'Q3')
    assert result.tables['Banking']['rows'][0][5] is None
    assert any('invalid/blank Cost' in i for i in result.issues)


def test_streamlit_renders_processed_results():
    from streamlit.testing.v1 import AppTest
    rev, hc = baseline(baseline_input(), 'Q3'), headcount(hc_input())
    app = AppTest.from_file('app.py')
    app.session_state['results'] = [(rev, 'Quarterly Baseline Updated.xlsx', 'rev'),
                                    (hc, 'Headcount Updated.xlsx', 'hc'),
                                    (margin(rev, hc), 'Margin Calculation.xlsx', 'gm')]
    app.run()
    assert not app.exception
    assert len(app.dataframe) == 13


def test_formula_values_from_verified_cache_are_used():
    raw = io.BytesIO()
    wb = xlsxwriter.Workbook(raw, {'in_memory': True})
    ws = wb.add_worksheet('Q3 Baseline Data')
    for c, label in {2: 'Segment', 3: 'LoB', 4: 'Service Line',
                     **dict(zip([7, 10, 18, 21], revenue_headers('Q3')))}.items():
        ws.write(0, c - 1, label)
    for c, value in zip([2, 3, 4], ['A', 'B', 'C']):
        ws.write(1, c - 1, value)
    for c in [7, 10, 18, 21]:
        ws.write_formula(1, c - 1, '=1+2', None, 3)
    wb.close()
    result = baseline(raw.getvalue(), 'Q3', formula_caches_verified=True)
    assert result.detail['Segment View'][('A',)] == [3, 3, 3, 3]
    output = openpyxl.load_workbook(io.BytesIO(result.workbook))
    assert output['Q3 Baseline Data']['G2'].value == '=1+2'


def test_unsupported_excel_features_rejected_before_mutation():
    import zipfile
    raw = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(baseline_input())) as source, zipfile.ZipFile(raw, 'w') as dest:
        for name in source.namelist():
            dest.writestr(name, source.read(name))
        dest.writestr('xl/externalLinks/externalLink1.xml', '<externalLink/>')
    with pytest.raises(ValidationError, match='Unsupported workbook features'):
        baseline(raw.getvalue(), 'Q3')


def test_formula_looking_identifiers_remain_literal_text():
    wb = openpyxl.load_workbook(io.BytesIO(baseline_input()))
    ws = wb['Q3 Baseline Data']
    ws['C2'] = '=Literal account'
    ws['C2'].data_type = 's'
    result = baseline(save(wb), 'Q3')
    out = openpyxl.load_workbook(io.BytesIO(result.workbook))
    account_cells = [c for row in out['Account View'] for c in row if c.value == '=Literal account']
    assert account_cells and account_cells[0].data_type == 's'
